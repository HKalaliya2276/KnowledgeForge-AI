from docx import Document as DocxDocument

from app.chunking.document import Document


class DocxLoader:

    @staticmethod
    def load(file_path):

        doc = DocxDocument(file_path)

        text = "\n".join(
            para.text
            for para in doc.paragraphs
        )

        return Document(
            text=text,
            metadata={
                "source": str(file_path),
                "file_name": file_path.split("/")[-1],
                "file_type": "docx"
            }
        )