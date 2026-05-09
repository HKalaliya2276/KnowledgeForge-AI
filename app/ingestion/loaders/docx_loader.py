from pathlib import Path
from docx import Document as DocxDocument

from app.ingestion.document import Document


class DOCXLoader:

    @staticmethod
    def load(file_path: str) -> Document:

        path = Path(file_path)

        doc = DocxDocument(file_path)

        text = []

        for para in doc.paragraphs:
            if para.text.strip():
                text.append(para.text)

        content = "\n".join(text)

        metadata = {
            "source": str(path),
            "file_name": path.name,
            "file_type": "docx"
        }

        return Document(
            content=content,
            metadata=metadata
        )